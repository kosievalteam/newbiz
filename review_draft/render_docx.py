"""ReviewOpinion → .docx 렌더러 (python-docx).

지면 구성은 샘플 검토의견서를 따른다:
  1. 협의신청사업 및 검토의견 개요 (3개 표 + 유사중복 표)
  2. 검토의견 (종합의견 ○/-/* 위계, 사업내용 박스 ⇩ 점검 박스, 개선의견)
  〈사업간 비교표〉
  (선택) 참고1. 협의요청서 원문
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .schema import AxisOpinion, Dash, ReviewOpinion

FONT = "맑은 고딕"
BASE_PT = 10
NOTE_PT = 9
GREY = RGBColor(0x40, 0x40, 0x40)


# ---------------------------------------------------------------------------
# 저수준 헬퍼
# ---------------------------------------------------------------------------

def _set_font(run, size=BASE_PT, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)
    if color is not None:
        run.font.color.rgb = color


def _para(doc_or_cell, text="", *, size=BASE_PT, bold=False, indent_cm=0.0, hanging_cm=0.0, align=None,
          space_before=2, space_after=2, color=None):
    p = doc_or_cell.add_paragraph()
    if text:
        _set_font(p.add_run(text), size=size, bold=bold, color=color)
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.3
    if indent_cm or hanging_cm:
        pf.left_indent = Cm(indent_cm + hanging_cm)
        pf.first_line_indent = Cm(-hanging_cm) if hanging_cm else None
    if align:
        p.alignment = align
    return p


def _shade(cell, hex_fill: str):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _cell_text(cell, text: str, *, bold=False, size=BASE_PT, align=None):
    cell.text = ""
    first = True
    for line in str(text).split("\n"):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        _set_font(p.add_run(line), size=size, bold=bold)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        if align:
            p.alignment = align


def _kv_table(doc, rows: list[tuple[str, str]], label_width_cm=3.2, total_width_cm=16.0):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in rows:
        r = t.add_row().cells
        r[0].width = Cm(label_width_cm)
        r[1].width = Cm(total_width_cm - label_width_cm)
        _cell_text(r[0], k, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(r[0], "EDEDED")
        _cell_text(r[1], v)
    return t


def _grid_table(doc, header: list[str], rows: list[list[str]], first_col_cm=2.4, total_width_cm=16.0, size=NOTE_PT):
    ncol = len(header)
    t = doc.add_table(rows=1, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    other = (total_width_cm - first_col_cm) / max(ncol - 1, 1)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.width = Cm(first_col_cm if j == 0 else other)
        _cell_text(c, h, bold=True, size=size, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(c, "E8EEF7")
    for row in rows:
        cells = t.add_row().cells
        for j in range(ncol):
            val = row[j] if j < len(row) else ""
            cells[j].width = Cm(first_col_cm if j == 0 else other)
            _cell_text(cells[j], val, bold=(j == 0), size=size, align=WD_ALIGN_PARAGRAPH.CENTER if j == 0 else None)
            if j == 0:
                _shade(cells[j], "F3F3F3")
    return t


# ---------------------------------------------------------------------------
# 위계 항목
# ---------------------------------------------------------------------------

def _circle(doc, text):
    return _para(doc, "○ " + text, bold=True, indent_cm=0.3, hanging_cm=0.5, space_before=6)


def _dash(doc, d: Dash, indent_cm=0.9):
    label = f"({d.label}) " if d.label else ""
    _para(doc, "- " + label + d.text, indent_cm=indent_cm, hanging_cm=0.4)
    for i, n in enumerate(d.notes):
        mark = "*" * (i + 1)
        _para(doc, f"{mark} {n}", size=NOTE_PT, indent_cm=indent_cm + 0.5, hanging_cm=0.45, color=GREY, space_before=0, space_after=0)


def _axis(doc, ax: AxisOpinion):
    _circle(doc, ax.headline)
    if ax.group_labels and ax.dash_groups:
        for label, group in zip(ax.group_labels, ax.dash_groups):
            _para(doc, label, bold=True, indent_cm=0.8, space_before=3)
            for d in group:
                _dash(doc, d, indent_cm=1.1)
    for d in ax.dashes:
        _dash(doc, d)
    for r in ax.remarks:
        _para(doc, "※ " + r, size=NOTE_PT, indent_cm=0.9, hanging_cm=0.45, color=GREY)


def _box(doc, lines: list[tuple[str, int, bool]], fill: str, title: str | None = None, center=False):
    """1×1 표 박스. lines = [(텍스트, 들여쓰기 단계, 굵게)]"""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    cell.width = Cm(15.2)
    _shade(cell, fill)
    cell.text = ""
    first = True
    if title:
        p = cell.paragraphs[0]
        first = False
        _set_font(p.add_run(title), bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
    for text, level, bold in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        _set_font(p.add_run(text), bold=bold, size=BASE_PT if level == 0 else NOTE_PT)
        p.paragraph_format.left_indent = Cm(0.2 + 0.6 * level)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return t


def _content_boxes(doc, op: ReviewOpinion):
    cb = op.content_box
    lines: list[tuple[str, int, bool]] = [("▪ (지원대상) " + cb.지원대상, 0, False), ("▪ (추진내용) " + cb.추진내용, 0, False)]
    lines += [("- " + s, 1, False) for s in cb.추진내용_세부]
    lines.append(("▪ (지원규모) " + cb.지원규모, 0, False))
    lines += [("- " + s, 1, False) for s in cb.지원규모_세부]
    lines.append(("▪ (수행기관) " + cb.수행기관, 0, False))
    _para(doc, "", space_before=1, space_after=1)
    _box(doc, lines, "F7F7F7", title="사업내용")
    _para(doc, "⇩", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1, space_after=1)
    _box(doc, [(l, 0, False) for l in op.check_box.lines], "FFF8E1", title="사업내용 점검", center=True)
    _para(doc, "", space_before=1, space_after=1)


# ---------------------------------------------------------------------------
# 문서 조립
# ---------------------------------------------------------------------------

def _setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BASE_PT)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def render_docx(op: ReviewOpinion, out_path: str | Path, *, appendix_text: str | None = None, draft_banner: bool = True) -> Path:
    doc = Document()
    _setup(doc)

    _para(doc, "중소기업지원사업 사전협의 검토의견서", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    if draft_banner:
        _para(doc, "※ 자동 생성 초안 — 검토자 확인 후 발송", size=NOTE_PT, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0xC0, 0, 0), space_after=8)

    # 1. 개요
    _para(doc, "1. 협의신청사업 및 검토의견 개요", size=12, bold=True, space_before=8)
    _para(doc, "□ 협의신청사업 개요", bold=True, space_before=4)
    ov = op.overview
    _kv_table(doc, [("1. 사업명", ov.사업명), ("2. 신청기관", ov.신청기관), ("3. 지원유형", ov.지원유형),
                    ("4. 특이사항", ov.특이사항), ("5. 검토자", ov.검토자 or "")])
    _para(doc, "□ 사업개요 추가정보", bold=True, space_before=8)
    ex = op.extra_info
    _kv_table(doc, [("1-1. 근거법령", ex.근거법령), ("1-2. 상위정책", ex.상위정책), ("1-3. 사업목적", ex.사업목적),
                    ("2-1. 지원대상", ex.지원대상), ("2-2. 지원내용", ex.지원내용), ("2-3. 지원규모", ex.지원규모),
                    ("2-4. 전달체계", ex.전달체계), ("3-1. 연계사업", ex.연계사업)])
    _para(doc, "□ 검토의견 요약", bold=True, space_before=8)
    s = op.summary
    _grid_table(doc, ["1. 종합의견", "2. 사업타당성", "3. 사업적합성", "4. 유사중복성"],
                [[s.종합의견, s.사업타당성, s.사업적합성, s.유사중복성]], first_col_cm=4.0, size=BASE_PT)
    _para(doc, "※ 유사중복검토 사업(유사중복사업이 있을 경우에만 작성)", size=NOTE_PT, space_before=6)
    rows = [[p.부처, p.사업명, p.검토의견] for p in op.similar_programs] or [["-", "해당 없음", "-"]]
    _grid_table(doc, ["부 처", "사업명", "검토의견"], rows, first_col_cm=2.4)

    # 2. 검토의견
    doc.add_page_break()
    _para(doc, "2. 검토의견", size=12, bold=True)
    _para(doc, "1. 종합의견", size=11, bold=True, space_before=4)
    _axis(doc, op.validity)
    _axis(doc, op.suitability)
    _content_boxes(doc, op)
    _axis(doc, op.duplication)

    _para(doc, "2. 개선의견", size=11, bold=True, space_before=10)
    for imp in op.improvements:
        _circle(doc, imp.headline)
        for d in imp.dashes:
            _dash(doc, d)

    # 비교표
    _para(doc, "〈 사업간 비교표 〉", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=14, space_after=4)
    _grid_table(doc, ["구 분"] + op.comparison.columns, op.comparison.rows, first_col_cm=2.0)

    if op.reviewer_notes:
        doc.add_page_break()
        _para(doc, "[검토자 확인 메모] (발송 전 삭제)", bold=True, color=RGBColor(0xC0, 0, 0))
        for n in op.reviewer_notes:
            _para(doc, "- " + n, indent_cm=0.3, hanging_cm=0.4, color=RGBColor(0xC0, 0, 0))

    if appendix_text:
        doc.add_page_break()
        _para(doc, "참고1. 협의요청서 (원문 추출)", size=12, bold=True)
        for line in appendix_text.splitlines():
            if line.strip():
                _para(doc, line, size=NOTE_PT, space_before=0, space_after=0)

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
