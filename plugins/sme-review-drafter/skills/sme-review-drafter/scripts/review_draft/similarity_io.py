"""유사도 분석 결과(similarity_check.py 산출물) 읽기.

- similarity_scores.xlsx : 시트 '전체'(또는 내역사업별 시트)의 상위 행
- top10.json             : {"weights":…, "model":…, "n_existing":…, "top": {unit: [row…]}}
둘 다 열 이름은 similarity_check.py 가 쓰는 그대로 사용한다.
"""

from __future__ import annotations

import json
from pathlib import Path

from .schema import SimilarityCandidate, SimilarityReport

_AXES = ["사업목적", "지원내용", "지원대상", "전달체계"]


def _row_to_candidate(r: dict, top3_note: bool) -> SimilarityCandidate:
    name = str(r.get("공고이름") or r.get("사업명") or "")
    name = name.replace("2026년 ", "").replace("2025년 ", "", 1) if name.startswith(("2026년 ", "2025년 ")) else name
    year = str(r.get("연도") or "")
    if year and year not in name:
        name = f"{name} ({year})"
    org = " · ".join(x for x in (str(r.get("부처") or ""), str(r.get("기관") or "")) if x)
    rank = int(r.get("순위") or 0)
    return SimilarityCandidate(
        순위=rank,
        사업명=name,
        부처_기관=org,
        단위=str(r.get("단위") or r.get("출처") or ""),
        종합유사도=float(r.get("종합유사도(0~100)") or r.get("종합유사도") or 0),
        사업목적=float(r.get("사업목적_유사도") or 0),
        지원내용=float(r.get("지원내용_유사도") or 0),
        지원대상=float(r.get("지원대상_유사도") or 0),
        전달체계=float(r.get("전달체계_유사도") or 0),
        비고="비교표 반영(상위 3)" if (top3_note and rank <= 3) else "",
    )


def load_similarity(path: str | Path, unit: str = "전체", top: int = 10, mark_top3: bool = True) -> SimilarityReport:
    p = Path(path)
    if p.suffix.lower() == ".json":
        data = json.loads(p.read_text(encoding="utf-8"))
        rows = data.get("top", {}).get(unit) or next(iter(data.get("top", {}).values()), [])
        n = data.get("n_existing")
        model = data.get("model", "")
        w = data.get("weights", {})
    else:
        import openpyxl

        wb = openpyxl.load_workbook(str(p), read_only=True)
        sheet = unit if unit in wb.sheetnames else next(s for s in wb.sheetnames if s != "가중치")
        ws = wb[sheet]
        it = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(it)]
        rows = []
        for vals in it:
            rows.append(dict(zip(header, vals)))
            if len(rows) >= top:
                break
        n = ws.max_row - 1 if ws.max_row else None
        model, w = "", {}
    cands = [_row_to_candidate(r, mark_top3) for r in rows[:top]]
    wtxt = "·".join(f"{k} {float(v):.2f}" for k, v in w.items()) if w else "사업목적 21.05·지원내용 31.58·지원대상 31.58·전달체계 15.79"
    method = f"기존사업 {n:,}개 단위" if n else "기존사업 단위"
    method += f", {model or 'BAAI/bge-m3'} 임베딩 코사인 유사도의 가중합(가중치 {wtxt}), 단위 '{unit}'"
    return SimilarityReport(method=method, candidates=cands, notes=[
        "임베딩 유사도는 서술문의 의미적 근접성을 측정한 선별 지표이며 중복 판정 결과가 아님. 상위 후보는 예산서·공고문 대조 후 비교표에 반영",
        "전달체계 축은 소관·수행기관·시행방법으로 근사하므로 동일 수행기관 사업의 유사도가 높게 산출되는 경향",
    ])


def report_to_prompt_text(rep: SimilarityReport) -> str:
    lines = ["# 유사도 분석 결과 (기존사업 상위 후보)", f"- 분석 기준: {rep.method}", "",
             "| 순위 | 기존사업 단위 | 부처·기관 | 종합 | 사업목적 | 지원내용 | 지원대상 | 전달체계 |", "|---|---|---|---|---|---|---|---|"]
    for c in rep.candidates:
        lines.append(f"| {c.순위} | {c.사업명} | {c.부처_기관} | {c.종합유사도:.1f} | {c.사업목적:.2f} | {c.지원내용:.2f} | {c.지원대상:.2f} | {c.전달체계:.2f} |")
    lines += ["", "상위 3개 후보를 유사·중복성 검토와 사업간 비교표의 기존사업으로 사용하고, 4~10위는 필요 시 연계·조정 검토 대상으로 언급한다."]
    return "\n".join(lines)
