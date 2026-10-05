import json
from pathlib import Path

from docx import Document

from review_draft.render_docx import render_docx
from review_draft.render_md import render_md
from review_draft.schema import ReviewOpinion, strict_schema

FIX = Path(__file__).parent / "fixtures" / "opinion_sample.json"


def _op():
    return ReviewOpinion.model_validate_json(FIX.read_text(encoding="utf-8"))


def test_schema_roundtrip():
    op = _op()
    again = ReviewOpinion.model_validate_json(op.model_dump_json())
    assert again == op


def test_strict_schema_flags():
    schema = strict_schema(ReviewOpinion)

    def check(node):
        if isinstance(node, dict):
            if node.get("type") == "object" and "properties" in node:
                assert node["additionalProperties"] is False
                assert set(node["required"]) == set(node["properties"])
                for v in node["properties"].values():
                    assert "default" not in v
            for v in node.values():
                check(v)
        elif isinstance(node, list):
            for v in node:
                check(v)

    check(schema)
    json.dumps(schema)  # 직렬화 가능


def test_render_md():
    md = render_md(_op())
    assert "○ (사업 타당성)" in md
    assert "* (법적근거)" in md
    assert "사업내용 점검" in md
    assert "| 구 분 |" in md
    assert "(중복수혜 방지)" in md


def test_render_docx(tmp_path):
    out = render_docx(_op(), tmp_path / "x.docx", appendix_text="원문 1행\n원문 2행")
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "○ (사업 적합성)" in text
    assert "⇩" in text
    assert "참고1. 협의요청서" in text and "원문 2행" in text
    tables = doc.tables
    assert any("사업내용 점검" in t.rows[0].cells[0].text for t in tables)
    cmp = [t for t in tables if t.rows[0].cells[0].text.strip() == "구 분"][0]
    assert len(cmp.columns) == 3
    assert cmp.rows[1].cells[0].text == "사업부서"
