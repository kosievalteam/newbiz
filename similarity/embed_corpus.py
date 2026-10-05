"""기존사업 단위 코퍼스를 bge-m3 로 임베딩해 .npz + .json 으로 저장한다 (Supabase 적재용).

입력(둘 중 하나)
  --xlsx similarity_scores.xlsx  : similarity_check.py 산출물의 '전체' 시트 (행·단위·연도·세부사업명·… + 4개 축 _텍스트 열)
  --units units.json             : [{"행":…, "사업목적_텍스트":…, …}] 형식의 JSON
출력
  <out>/corpus_embeddings.npz    : purpose/content/target/delivery (N×1024, float32, L2 정규화)
  <out>/corpus_meta.json         : 단위 메타데이터 + 축별 텍스트 (적재 시 함께 올림)
사용
  python similarity/embed_corpus.py --xlsx similarity/results/유사도_산출결과_청년금융혁신.xlsx --out output/embeddings
"""

from __future__ import annotations

import argparse
import json
import time
from pathlib import Path

AXES = {"purpose": "사업목적_텍스트", "content": "지원내용_텍스트", "target": "지원대상_텍스트", "delivery": "전달체계_텍스트"}
META_COLS = ["행", "단위", "연도", "부처", "기관", "공고이름", "세부사업명", "내역사업명", "내내역명", "시행방법", "규모", "연계공고", "대분류", "중분류"]


def load_units(xlsx: str | None, units_json: str | None) -> list[dict]:
    if units_json:
        return json.loads(Path(units_json).read_text(encoding="utf-8"))
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ws = wb["전체"] if "전체" in wb.sheetnames else wb.worksheets[1]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h) for h in next(it)]
    rows = [dict(zip(hdr, r)) for r in it]
    return [r for r in rows if r.get("행")]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx")
    ap.add_argument("--units")
    ap.add_argument("--out", default="output/embeddings")
    ap.add_argument("--model", default="BAAI/bge-m3")
    ap.add_argument("--batch", type=int, default=16)
    ap.add_argument("--limit", type=int, default=0, help="테스트용: 앞 N개 단위만")
    ap.add_argument("--max-seq", type=int, default=512, help="토큰 길이 상한 (bge-m3 기본 8192; 512 면 CPU 에서 수 배 빠르고 평균 길이의 텍스트는 잘리지 않음)")
    ap.add_argument("--threads", type=int, default=0)
    a = ap.parse_args()
    if not (a.xlsx or a.units):
        raise SystemExit("--xlsx 또는 --units 가 필요합니다")

    import numpy as np
    from sentence_transformers import SentenceTransformer

    units = load_units(a.xlsx, a.units)
    if a.limit:
        units = units[: a.limit]
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    meta = []
    for u in units:
        m = {k: ("" if u.get(k) is None else str(u.get(k))) for k in META_COLS if k in u}
        for ax, col in AXES.items():
            m[ax] = "" if u.get(col) is None else str(u.get(col))
        meta.append(m)
    (out / "corpus_meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")

    if a.threads:
        import torch

        torch.set_num_threads(a.threads)
    model = SentenceTransformer(a.model, device="cpu")
    model.max_seq_length = a.max_seq
    arrays = {}
    for ax in AXES:
        texts = [m[ax] or " " for m in meta]
        t0 = time.time()
        chunks = []
        step = 256
        for i in range(0, len(texts), step):
            chunks.append(model.encode(texts[i : i + step], normalize_embeddings=True, batch_size=a.batch, show_progress_bar=False))
            print(f"{ax}: {min(i + step, len(texts))}/{len(texts)} {time.time() - t0:.0f}s", flush=True)
        arrays[ax] = np.concatenate(chunks).astype("float32")
        np.save(out / f"partial_{ax}.npy", arrays[ax])
        print(f"{ax}: 완료 {time.time() - t0:.0f}s", flush=True)
    np.savez_compressed(out / "corpus_embeddings.npz", **arrays)
    print(f"저장: {out / 'corpus_embeddings.npz'} ({len(meta)} units, dim {arrays['purpose'].shape[1]})")


if __name__ == "__main__":
    main()
